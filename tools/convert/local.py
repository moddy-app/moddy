"""Local (in-process) converters of /convert — everything except ffmpeg and the cloud.

``convert_bytes(src, dst, raw, filename)`` is synchronous and Discord-free: the
service runs it in ``asyncio.to_thread``. Every converter

* validates sizes **before** decoding (input caps, spreadsheet cell cap, archive
  entry/size/ratio caps, PDF page caps) and raises :class:`MediaError` with a
  stable code — never a library exception;
* imports its heavy library lazily, so ``import tools.convert.local`` is free;
* never executes anything from the file (``yaml.safe_load`` only, no formulas
  evaluated, archive member names sanitised, links and devices skipped).
"""

from __future__ import annotations

import csv
import io
import json
import re
import tarfile
import tempfile
import threading
import zipfile
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath
from typing import Any, List, Optional, Sequence, Tuple

from .. import image_ops
from ..media_errors import MediaError
from . import formats as F


@dataclass
class ConvertOutput:
    """What a conversion produced: one file (``name`` + ``data``)."""

    name: str
    data: bytes
    format_key: str
    notes: List[str] = field(default_factory=list)   # i18n keys under tools.convert.notes.*
    pages: int = 0                                   # PDF -> images: how many pages are inside
    width: int = 0                                   # images: result size
    height: int = 0

    @property
    def mime(self) -> str:
        return F.FORMATS[self.format_key].mime if self.format_key in F.FORMATS else "application/octet-stream"


# --------------------------------------------------------------------------- #
# Text helpers
# --------------------------------------------------------------------------- #


def decode_text(raw: bytes) -> str:
    """Bytes -> str: BOMs first, then UTF-8, then the legacy Windows code page."""
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8", "replace")
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return raw.decode("utf-16")
        except UnicodeDecodeError:
            pass
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("cp1252", "replace")


def _text_input(raw: bytes) -> str:
    if len(raw) > F.MAX_TEXT_BYTES:
        raise MediaError("too_large", max=F.MAX_TEXT_BYTES // (1024 * 1024))
    return decode_text(raw)


def _check_zip_bomb(raw: bytes) -> None:
    """Refuse a zip-based document (docx/xlsx/ods) that would expand absurdly."""
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            infos = zf.infolist()
            total = sum(i.file_size for i in infos)
            if len(infos) > 5_000 or total > F.MAX_ARCHIVE_UNCOMPRESSED:
                raise MediaError("archive_too_big", max=F.MAX_ARCHIVE_UNCOMPRESSED // (1024 * 1024))
            for i in infos:
                if i.compress_size and i.file_size / i.compress_size > F.MAX_ARCHIVE_RATIO \
                        and i.file_size > 1024 * 1024:
                    raise MediaError("unsafe_archive")
    except zipfile.BadZipFile as exc:
        raise MediaError("unreadable") from exc


# --------------------------------------------------------------------------- #
# Tables (CSV / XLSX / ODS) and structured data (JSON / YAML / TOML / XML)
# --------------------------------------------------------------------------- #

Rows = List[List[Any]]


def _check_cells(rows: Rows) -> None:
    cells = sum(len(r) for r in rows)
    if cells > F.MAX_SHEET_CELLS:
        raise MediaError("too_many_cells", max=F.MAX_SHEET_CELLS)


def read_csv_rows(text: str) -> Rows:
    sample = text[:4096]
    delimiter = ","
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        counts = {d: sample.count(d) for d in (",", ";", "\t")}
        delimiter = max(counts, key=counts.get) if any(counts.values()) else ","
    csv.field_size_limit(1 << 20)
    rows: Rows = []
    cells = 0
    try:
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter):
            cells += len(row)
            if cells > F.MAX_SHEET_CELLS:
                raise MediaError("too_many_cells", max=F.MAX_SHEET_CELLS)
            rows.append(row)
    except csv.Error as exc:
        raise MediaError("unreadable") from exc
    return rows


def write_csv_bytes(rows: Rows, *, bom: bool = False) -> bytes:
    out = io.StringIO(newline="")
    writer = csv.writer(out, lineterminator="\r\n" if bom else "\n")
    for row in rows:
        writer.writerow(["" if v is None else v for v in row])
    return (("﻿" if bom else "") + out.getvalue()).encode("utf-8")


def read_xlsx_rows(raw: bytes) -> Tuple[Rows, List[str]]:
    _check_zip_bomb(raw)
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise MediaError("unreadable") from exc
    notes: List[str] = []
    try:
        if len(wb.sheetnames) > 1:
            notes.append("first_sheet_only")
        ws = wb[wb.sheetnames[0]]
        rows: Rows = []
        cells = 0
        for row in ws.iter_rows(values_only=True):
            values = list(row)
            while values and values[-1] is None:
                values.pop()
            cells += len(values)
            if cells > F.MAX_SHEET_CELLS:
                raise MediaError("too_many_cells", max=F.MAX_SHEET_CELLS)
            rows.append(values)
        while rows and not rows[-1]:
            rows.pop()
        return rows, notes
    except MediaError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise MediaError("unreadable") from exc
    finally:
        wb.close()


def write_xlsx_bytes(rows: Rows) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Sheet1")
    for row in rows:
        cells = []
        for value in row:
            if isinstance(value, str) and value[:1] in ("=", "+", "-", "@") and not _is_number(value):
                # Text typed by a human must stay text, never a formula.
                cell = WriteOnlyCell(ws, value=value)
                cell.data_type = "s"
                cells.append(cell)
            elif isinstance(value, (dict, list)):
                cells.append(json.dumps(value, ensure_ascii=False))
            else:
                cells.append(value)
        ws.append(cells)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _is_number(text: str) -> bool:
    try:
        float(text)
        return True
    except ValueError:
        return False


_ODS_NS = {
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}


def read_ods_rows(raw: bytes) -> Tuple[Rows, List[str]]:
    """First sheet of an ODS file (stdlib only: no odfpy)."""
    import xml.etree.ElementTree as ET
    _check_zip_bomb(raw)
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            content = zf.read("content.xml")
    except (KeyError, zipfile.BadZipFile) as exc:
        raise MediaError("unreadable") from exc

    def q(prefix: str, name: str) -> str:
        return "{%s}%s" % (_ODS_NS[prefix], name)

    rows: Rows = []
    notes: List[str] = []
    cells = 0
    tables_seen = 0
    try:
        for _event, el in ET.iterparse(io.BytesIO(content), events=("end",)):
            if el.tag != q("table", "table"):
                continue
            tables_seen += 1
            if tables_seen > 1:
                notes.append("first_sheet_only")
                break
            for tr in el.iter(q("table", "table-row")):
                repeat = int(tr.get(q("table", "number-rows-repeated"), "1") or 1)
                row: List[Any] = []
                for tc in tr:
                    if tc.tag not in (q("table", "table-cell"), q("table", "covered-table-cell")):
                        continue
                    n = int(tc.get(q("table", "number-columns-repeated"), "1") or 1)
                    value = _ods_value(tc, q)
                    row.extend([value] * min(n, 1024 if value is None else n))
                    if len(row) > 16_384:
                        raise MediaError("too_many_cells", max=F.MAX_SHEET_CELLS)
                while row and row[-1] is None:
                    row.pop()
                for _ in range(min(repeat, 1 if not row else 1000)):
                    rows.append(list(row))
                    cells += len(row)
                if cells > F.MAX_SHEET_CELLS:
                    raise MediaError("too_many_cells", max=F.MAX_SHEET_CELLS)
            el.clear()
    except MediaError:
        raise
    except Exception as exc:  # noqa: BLE001 — ET.ParseError, ValueError...
        raise MediaError("unreadable") from exc
    while rows and not rows[-1]:
        rows.pop()
    return rows, notes


def _ods_value(cell, q) -> Any:
    vtype = cell.get(q("office", "value-type"))
    if vtype is None:
        return None
    if vtype in ("float", "percentage", "currency"):
        try:
            number = float(cell.get(q("office", "value"), "0"))
            return int(number) if number.is_integer() else number
        except ValueError:
            return None
    if vtype == "boolean":
        return cell.get(q("office", "boolean-value")) == "true"
    if vtype in ("date", "time"):
        return cell.get(q("office", "date-value")) or cell.get(q("office", "time-value"))
    parts = ["".join(p.itertext()) for p in cell.findall(q("text", "p"))]
    return "\n".join(parts) if parts else ""


def write_ods_bytes(rows: Rows) -> bytes:
    from xml.sax.saxutils import escape
    out = io.StringIO()
    out.write(
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="%(office)s" xmlns:table="%(table)s" '
        'xmlns:text="%(text)s" office:version="1.2"><office:body><office:spreadsheet>'
        '<table:table table:name="Sheet1">' % _ODS_NS)
    for row in rows:
        out.write("<table:table-row>")
        for value in row:
            if value is None or value == "":
                out.write("<table:table-cell/>")
            elif isinstance(value, bool):
                out.write('<table:table-cell office:value-type="boolean" office:boolean-value="%s">'
                          "<text:p>%s</text:p></table:table-cell>" % (str(value).lower(), str(value).upper()))
            elif isinstance(value, (int, float)):
                out.write('<table:table-cell office:value-type="float" office:value="%s">'
                          "<text:p>%s</text:p></table:table-cell>" % (value, value))
            else:
                text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, default=str)
                text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
                out.write('<table:table-cell office:value-type="string"><text:p>%s</text:p>'
                          "</table:table-cell>" % escape(text))
        out.write("</table:table-row>")
    out.write("</table:table></office:spreadsheet></office:body></office:document-content>")
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" '
        'manifest:version="1.2">'
        '<manifest:file-entry manifest:full-path="/" manifest:media-type='
        '"application/vnd.oasis.opendocument.spreadsheet"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.spreadsheet")
        zf.writestr("content.xml", out.getvalue())
        zf.writestr("META-INF/manifest.xml", manifest)
    return buf.getvalue()


def _dedupe_headers(header: Sequence[Any]) -> List[str]:
    seen = {}
    names: List[str] = []
    for i, value in enumerate(header, 1):
        name = str(value).strip() if value not in (None, "") else f"column_{i}"
        count = seen.get(name, 0)
        seen[name] = count + 1
        names.append(name if count == 0 else f"{name}_{count + 1}")
    return names


def rows_to_records(rows: Rows) -> List[dict]:
    """First row = header, the rest = one dict per row."""
    if not rows:
        return []
    header = _dedupe_headers(rows[0])
    records = []
    for row in rows[1:]:
        records.append({header[i]: (row[i] if i < len(row) else None) for i in range(len(header))})
    return records


def records_to_rows(obj: Any) -> Rows:
    """A list of flat-ish objects (or a single object) -> header + rows."""
    if isinstance(obj, dict):
        if len(obj) == 1 and isinstance(next(iter(obj.values())), list):
            obj = next(iter(obj.values()))      # {"items": [...]} (what TOML output wraps)
        else:
            obj = [obj]
    if not isinstance(obj, list) or not obj:
        raise MediaError("not_tabular")
    if all(isinstance(r, list) for r in obj):
        rows = [list(r) for r in obj]
    elif all(isinstance(r, dict) for r in obj):
        columns: List[str] = []
        for record in obj:
            for key in record:
                if key not in columns:
                    columns.append(str(key))
        rows = [columns] + [[_cell(record.get(c)) for c in columns] for record in obj]
    else:
        rows = [["value"]] + [[_cell(v)] for v in obj]
    _check_cells(rows)
    return rows


def _cell(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return value


_ALIAS = re.compile(r"(?m)(?<![\w'\"])\*[A-Za-z0-9_-]+")


def parse_data(key: str, raw: bytes) -> Any:
    text = _text_input(raw)
    try:
        if key == "json":
            return json.loads(text)
        if key == "yaml":
            if len(_ALIAS.findall(text)) > 100:   # alias expansion ("billion laughs") guard
                raise MediaError("unsafe_data")
            import yaml
            return yaml.safe_load(text)
        if key == "toml":
            import tomllib
            return tomllib.loads(text)
        if key == "xml":
            import xmltodict
            return xmltodict.parse(text)
        if key == "csv":
            return rows_to_records(read_csv_rows(text))
    except MediaError:
        raise
    except Exception as exc:  # noqa: BLE001 — each parser has its own error types
        raise MediaError("invalid_data", detail=str(exc).splitlines()[0][:120] if str(exc) else "") from exc
    raise MediaError("unsupported")


def _drop_none(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(k): _drop_none(v) for k, v in value.items() if v is not None}
    if isinstance(value, (list, tuple)):
        return [_drop_none(v) for v in value if v is not None]
    return value


def dump_data(key: str, obj: Any, notes: List[str]) -> bytes:
    try:
        if key == "json":
            return json.dumps(obj, indent=2, ensure_ascii=False, default=str).encode("utf-8")
        if key == "yaml":
            import yaml
            return yaml.safe_dump(json.loads(json.dumps(obj, default=str)), allow_unicode=True,
                                  sort_keys=False).encode("utf-8")
        if key == "toml":
            import tomli_w
            clean = _drop_none(obj)
            if clean != obj:
                notes.append("null_dropped")
            if isinstance(clean, list):
                clean = {"items": clean}
            if not isinstance(clean, dict):
                raise MediaError("unconvertible")
            return tomli_w.dumps(clean).encode("utf-8")
        if key == "xml":
            import xmltodict
            if isinstance(obj, dict) and len(obj) == 1 and not isinstance(next(iter(obj.values())), list):
                doc = obj
            elif isinstance(obj, list):
                doc = {"root": {"item": obj}}
            else:
                doc = {"root": obj}
            return xmltodict.unparse(json.loads(json.dumps(doc, default=str)), pretty=True,
                                     full_document=True).encode("utf-8")
        if key == "csv":
            return write_csv_bytes(records_to_rows(obj))
    except MediaError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise MediaError("unconvertible") from exc
    raise MediaError("unsupported")


def convert_data(src: str, dst: str, raw: bytes) -> Tuple[bytes, List[str]]:
    notes: List[str] = []
    obj = parse_data(src, raw)
    if obj is None:
        raise MediaError("empty_file")
    return dump_data(dst, obj, notes), notes


def read_rows(src: str, raw: bytes) -> Tuple[Rows, List[str]]:
    if src == "csv":
        return read_csv_rows(_text_input(raw)), []
    if src == "xlsx":
        return read_xlsx_rows(raw)
    if src == "ods":
        return read_ods_rows(raw)
    if src == "json":
        return records_to_rows(parse_data("json", raw)), []
    raise MediaError("unsupported")


def convert_sheet(src: str, dst: str, raw: bytes) -> Tuple[bytes, List[str]]:
    rows, notes = read_rows(src, raw)
    if not rows:
        raise MediaError("empty_file")
    _check_cells(rows)
    if dst == "csv":
        return write_csv_bytes(rows, bom=src in ("xlsx", "ods")), notes
    if dst == "xlsx":
        return write_xlsx_bytes(rows), notes
    if dst == "ods":
        return write_ods_bytes(rows), notes
    if dst == "json":
        return json.dumps(rows_to_records(rows), indent=2, ensure_ascii=False,
                          default=str).encode("utf-8"), notes
    raise MediaError("unsupported")


# --------------------------------------------------------------------------- #
# Documents: Markdown / HTML / TXT / DOCX
# --------------------------------------------------------------------------- #

_HTML_SHELL = ('<!doctype html>\n<html>\n<head>\n<meta charset="utf-8">\n'
               '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
               "<title>{title}</title>\n</head>\n<body>\n{body}\n</body>\n</html>\n")


def _soup(html: str):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()
    return soup


def html_to_text(html: str) -> str:
    text = _soup(html).get_text("\n")
    lines = [line.strip() for line in text.splitlines()]
    text = "\n".join(lines)
    return re.sub(r"\n{3,}", "\n\n", text).strip() + "\n"


def html_to_markdown(html: str) -> str:
    from markdownify import markdownify
    out = markdownify(str(_soup(html)), heading_style="ATX", bullets="-")
    return re.sub(r"\n{3,}", "\n\n", out).strip() + "\n"


def markdown_to_html(text: str) -> str:
    import markdown
    return markdown.markdown(text, extensions=["extra", "sane_lists"])


def docx_to_html(raw: bytes, notes: List[str]) -> str:
    _check_zip_bomb(raw)
    import mammoth
    try:
        result = mammoth.convert_to_html(
            io.BytesIO(raw),
            convert_image=mammoth.images.img_element(lambda image: {"alt": "(image)"}))
    except Exception as exc:  # noqa: BLE001
        raise MediaError("unreadable") from exc
    if "<img" in result.value:
        notes.append("images_dropped")
    return result.value


def convert_doc(src: str, dst: str, raw: bytes, filename: str) -> Tuple[bytes, List[str]]:
    notes: List[str] = []
    title = re.sub(r"[<>&]", "", Path(filename).stem or "document")
    if src == "docx":
        if len(raw) > F.MAX_INPUT_BYTES:
            raise MediaError("too_large", max=F.MAX_INPUT_BYTES // (1024 * 1024))
        if dst == "txt":
            _check_zip_bomb(raw)
            import mammoth
            try:
                text = mammoth.extract_raw_text(io.BytesIO(raw)).value
            except Exception as exc:  # noqa: BLE001
                raise MediaError("unreadable") from exc
            return (text.strip() + "\n").encode("utf-8"), notes
        html = docx_to_html(raw, notes)
        if dst == "html":
            return _HTML_SHELL.format(title=title, body=html).encode("utf-8"), notes
        if dst == "md":
            return html_to_markdown(html).encode("utf-8"), notes
    text = _text_input(raw)
    if src == "md" and dst == "html":
        return _HTML_SHELL.format(title=title, body=markdown_to_html(text)).encode("utf-8"), notes
    if src == "md" and dst == "txt":
        return html_to_text(markdown_to_html(text)).encode("utf-8"), notes
    if src == "html" and dst == "md":
        return html_to_markdown(text).encode("utf-8"), notes
    if src == "html" and dst == "txt":
        return html_to_text(text).encode("utf-8"), notes
    if src == "txt" and dst == "txt":
        return text.replace("\r\n", "\n").encode("utf-8"), notes
    raise MediaError("unsupported")


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #

# pdfium is not thread-safe: one document at a time in the whole process.
_PDFIUM_LOCK = threading.Lock()


def pdf_to_text(raw: bytes) -> Tuple[bytes, List[str]]:
    notes: List[str] = []
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise MediaError("protected")
            except MediaError:
                raise
            except Exception as exc:  # noqa: BLE001
                raise MediaError("protected") from exc
        total = len(reader.pages)
        parts = []
        for page in reader.pages[: F.MAX_PDF_TEXT_PAGES]:
            parts.append((page.extract_text() or "").strip())
    except MediaError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise MediaError("unreadable") from exc
    if total > F.MAX_PDF_TEXT_PAGES:
        notes.append("pages_capped")
    text = "\n\n".join(p for p in parts if p).strip()
    if not text:
        raise MediaError("no_text")
    return (text + "\n").encode("utf-8"), notes


def pdf_to_images(raw: bytes, dst: str, stem: str) -> Tuple[bytes, str, List[str], int, Tuple[int, int]]:
    """Render up to 10 pages. Returns ``(data, "png"|"jpg"|"zip", notes, pages, size)``.

    One page -> the image itself; several -> a zip of ``<stem>-<n>.<ext>``.
    """
    import pypdfium2 as pdfium
    notes: List[str] = []
    images: List[Tuple[str, bytes]] = []
    size = (0, 0)
    with _PDFIUM_LOCK:
        try:
            doc = pdfium.PdfDocument(raw)
        except Exception as exc:  # noqa: BLE001
            msg = str(exc).lower()
            raise MediaError("protected" if "password" in msg else "unreadable") from exc
        try:
            total = len(doc)
            if total == 0:
                raise MediaError("unreadable")
            if total > F.MAX_PDF_RENDER_PAGES:
                notes.append("pages_capped")
            for index in range(min(total, F.MAX_PDF_RENDER_PAGES)):
                page = doc[index]
                try:
                    width, height = page.get_size()
                    scale = min(150 / 72, F.MAX_PDF_RENDER_SIDE / max(width, height, 1))
                    pil = page.render(scale=scale).to_pil().convert("RGB")
                finally:
                    page.close()
                size = size if index else pil.size
                buf = io.BytesIO()
                if dst == "jpg":
                    pil.save(buf, "JPEG", quality=88, optimize=True)
                else:
                    pil.save(buf, "PNG", optimize=False, compress_level=6)
                images.append((f"{stem}-{index + 1}.{dst}", buf.getvalue()))
        finally:
            doc.close()
    if len(images) == 1:
        return images[0][1], dst, notes, 1, size
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_STORED) as zf:
        for name, data in images:
            zf.writestr(name, data)
    return zbuf.getvalue(), "zip", notes, len(images), size


# --------------------------------------------------------------------------- #
# Archives
# --------------------------------------------------------------------------- #


def safe_member_name(name: str) -> Optional[str]:
    """Sanitised relative POSIX path, or ``None`` for anything that escapes."""
    name = name.replace("\\", "/")
    if not name or name.startswith("/") or re.match(r"^[A-Za-z]:", name):
        return None
    parts = [p for p in PurePosixPath(name).parts if p not in ("", ".")]
    if not parts or any(p == ".." for p in parts) or "\x00" in name:
        return None
    return "/".join(parts)


class _Budget:
    def __init__(self) -> None:
        self.files = 0
        self.bytes = 0

    def file(self, size: int = 0) -> None:
        self.files += 1
        self.bytes += size
        if self.files > F.MAX_ARCHIVE_FILES:
            raise MediaError("too_many_files", max=F.MAX_ARCHIVE_FILES)
        if self.bytes > F.MAX_ARCHIVE_UNCOMPRESSED:
            raise MediaError("archive_too_big", max=F.MAX_ARCHIVE_UNCOMPRESSED // (1024 * 1024))


def _copy_capped(src, dest: Path, budget: _Budget, limit_hint: int) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    written = 0
    with open(dest, "wb") as out:
        while True:
            chunk = src.read(1 << 20)
            if not chunk:
                break
            written += len(chunk)
            budget.bytes += len(chunk)
            if budget.bytes > F.MAX_ARCHIVE_UNCOMPRESSED:
                raise MediaError("archive_too_big", max=F.MAX_ARCHIVE_UNCOMPRESSED // (1024 * 1024))
            out.write(chunk)


def _extract_zip(raw: bytes, root: Path, budget: _Budget) -> List[str]:
    notes: List[str] = []
    try:
        zf = zipfile.ZipFile(io.BytesIO(raw))
    except zipfile.BadZipFile as exc:
        raise MediaError("unreadable") from exc
    with zf:
        for info in zf.infolist():
            if info.is_dir():
                continue
            if info.flag_bits & 0x1:
                raise MediaError("protected")
            name = safe_member_name(info.filename)
            if name is None or (info.external_attr >> 16) & 0o170000 == 0o120000:
                notes.append("entries_skipped")
                continue
            if info.compress_size and info.file_size / max(1, info.compress_size) > F.MAX_ARCHIVE_RATIO \
                    and info.file_size > 1024 * 1024:
                raise MediaError("unsafe_archive")
            budget.file()
            try:
                with zf.open(info) as src:
                    _copy_capped(src, root / name, budget, info.file_size)
            except (RuntimeError, zipfile.BadZipFile, NotImplementedError) as exc:
                raise MediaError("unreadable") from exc
    return notes


def _extract_tar(raw: bytes, root: Path, budget: _Budget) -> List[str]:
    notes: List[str] = []
    try:
        tf = tarfile.open(fileobj=io.BytesIO(raw), mode="r:*")
    except (tarfile.TarError, OSError, EOFError) as exc:
        raise MediaError("unreadable") from exc
    with tf:
        try:
            for member in tf:
                if not member.isfile():
                    if not member.isdir():
                        notes.append("entries_skipped")
                    continue
                name = safe_member_name(member.name)
                if name is None:
                    notes.append("entries_skipped")
                    continue
                budget.file()
                src = tf.extractfile(member)
                if src is None:
                    continue
                with src:
                    _copy_capped(src, root / name, budget, member.size)
        except (tarfile.TarError, OSError, EOFError) as exc:
            raise MediaError("unreadable") from exc
    return notes


def _extract_7z(raw: bytes, root: Path, budget: _Budget) -> List[str]:
    notes: List[str] = []
    try:
        import py7zr
        with py7zr.SevenZipFile(io.BytesIO(raw), mode="r") as archive:
            if archive.needs_password():
                raise MediaError("protected")
            entries = [e for e in archive.list() if not e.is_directory]
            if len(entries) > F.MAX_ARCHIVE_FILES:
                raise MediaError("too_many_files", max=F.MAX_ARCHIVE_FILES)
            if sum(e.uncompressed or 0 for e in entries) > F.MAX_ARCHIVE_UNCOMPRESSED:
                raise MediaError("archive_too_big", max=F.MAX_ARCHIVE_UNCOMPRESSED // (1024 * 1024))
            bad = [e.filename for e in entries if safe_member_name(e.filename) is None or e.is_symlink]
            if bad:
                notes.append("entries_skipped")
            wanted = [e.filename for e in entries if e.filename not in bad]
            archive.extract(path=root, targets=wanted)
    except MediaError:
        raise
    except Exception as exc:  # noqa: BLE001 — py7zr raises many types
        msg = str(exc).lower()
        raise MediaError("protected" if "password" in msg else "unreadable") from exc
    return notes


def _walk_files(root: Path) -> List[Path]:
    files = []
    for path in sorted(root.rglob("*")):
        if path.is_symlink() or not path.is_file():
            continue
        try:
            path.resolve().relative_to(root.resolve())
        except ValueError:
            continue
        files.append(path)
    return files


def convert_archive(src: str, dst: str, raw: bytes) -> Tuple[bytes, List[str]]:
    budget = _Budget()
    with tempfile.TemporaryDirectory(prefix="moddy-arch-") as tmp:
        root = Path(tmp) / "x"
        root.mkdir()
        if src == "zip":
            notes = _extract_zip(raw, root, budget)
        elif src == "tgz":
            notes = _extract_tar(raw, root, budget)
        else:
            notes = _extract_7z(raw, root, budget)
        files = _walk_files(root)
        if len(files) > F.MAX_ARCHIVE_FILES:
            raise MediaError("too_many_files", max=F.MAX_ARCHIVE_FILES)
        if not files:
            raise MediaError("empty_file")
        out = io.BytesIO()
        if dst == "zip":
            with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
                for path in files:
                    zf.write(path, path.relative_to(root).as_posix())
        elif dst == "tgz":
            with tarfile.open(fileobj=out, mode="w:gz") as tf:
                for path in files:
                    tf.add(path, arcname=path.relative_to(root).as_posix(), recursive=False)
        else:
            import py7zr
            with py7zr.SevenZipFile(out, "w") as archive:
                for path in files:
                    archive.write(path, arcname=path.relative_to(root).as_posix())
        return out.getvalue(), sorted(set(notes))


# --------------------------------------------------------------------------- #
# Entry point
# --------------------------------------------------------------------------- #


def convert_bytes(src: str, dst: str, raw: bytes, filename: str = "file") -> ConvertOutput:
    """Run the local conversion ``src -> dst``. Raises :class:`MediaError`."""
    route = F.route(src, dst)
    if route is None or route.cloud or route.engine == F.E_FFMPEG:
        raise MediaError("unsupported")
    if not raw:
        raise MediaError("empty_file")
    if len(raw) > F.MAX_INPUT_BYTES:
        raise MediaError("too_large", max=F.MAX_INPUT_BYTES // (1024 * 1024))
    name = F.output_name(filename, dst)
    stem = name.rsplit(".", 1)[0]
    engine = route.engine

    if engine == F.E_IMAGE:
        done = image_ops.convert(raw, dst)
        return ConvertOutput(name, done.data, dst, list(done.notes), width=done.width,
                             height=done.height)
    if engine == F.E_DATA:
        data, notes = convert_data(src, dst, raw)
        return ConvertOutput(name, data, dst, notes)
    if engine == F.E_SHEET:
        data, notes = convert_sheet(src, dst, raw)
        return ConvertOutput(name, data, dst, notes)
    if engine == F.E_DOC:
        data, notes = convert_doc(src, dst, raw, filename)
        return ConvertOutput(name, data, dst, notes)
    if engine == F.E_PDF:
        if dst == "txt":
            data, notes = pdf_to_text(raw)
            return ConvertOutput(name, data, "txt", notes)
        data, key, notes, pages, size = pdf_to_images(raw, dst, stem)
        out_name = name if key == dst else f"{stem}.zip"
        return ConvertOutput(out_name, data, key, notes, pages=pages, width=size[0], height=size[1])
    if engine == F.E_ARCHIVE:
        data, notes = convert_archive(src, dst, raw)
        return ConvertOutput(name, data, dst, notes)
    raise MediaError("unsupported")
