"""/image and /convert (docs/features/TOOLS_MEDIA.md).

Pure logic, converters, the background-removal lifecycle (with a fake ONNX
session: no model is downloaded), the CloudConvert adapter (fake HTTP), the
service's error mapping, the cards and the i18n parity of the 5 locales. No
network is used anywhere.
"""

from __future__ import annotations

import asyncio
import io
import json
import re
import types
import zipfile
from pathlib import Path

import discord
import pytest
from discord import ui
from PIL import Image

from gateway.adapters import cloudconvert as cc
from gateway.errors import (
    APIUnavailableError, ConfigurationError, ModelRateLimitError, ProviderError,
    QuotaExceededError, RateLimitError,
)
from gateway.spec import QuotaTarget
from services import media_tools_service as svc
from tools import bg_removal, image_ops
from tools.convert import detect, formats, local, media
from tools.media_errors import MediaError
from utils import tool_card as tc
from utils import tools_media_views as views
from utils.i18n import i18n

ROOT = Path(__file__).resolve().parent.parent
LOCALES = ("fr", "en-US", "es-ES", "pt-BR", "de")


def png(size=(64, 48), color=(200, 30, 30), mode="RGB") -> bytes:
    buf = io.BytesIO()
    Image.new(mode, size, color).save(buf, "PNG")
    return buf.getvalue()


def walk_text(view) -> str:
    out = []
    for item in view.walk_children():
        content = getattr(item, "content", None)
        if isinstance(content, str):
            out.append(content)
        label = getattr(item, "label", None)
        if isinstance(label, str):
            out.append(label)
    return "\n".join(out)


def custom_ids(view):
    return [getattr(i, "custom_id", None) for i in view.walk_children()
            if getattr(i, "custom_id", None)]


# --------------------------------------------------------------------------- #
# Format table and detection
# --------------------------------------------------------------------------- #


def test_every_conversion_target_is_a_known_format_and_fits_a_select():
    for src, row in formats.CONVERSIONS.items():
        assert src in formats.FORMATS
        assert len(row) <= 60
        for dst in row:
            assert dst in formats.FORMATS and (dst != src or src == "txt")  # txt -> txt re-encodes to UTF-8
        assert len(formats.targets(src, formats.Capabilities(cloud=True))) <= 25


def test_cloud_formats_are_hidden_without_a_key():
    assert "pdf" not in formats.targets("docx", formats.Capabilities(cloud=False))
    assert "pdf" in formats.targets("docx", formats.Capabilities(cloud=True))
    assert "epub" not in formats.targets("pdf", formats.Capabilities(cloud=False))
    # layout office documents are the only cloud route besides ebooks
    for src, row in formats.CONVERSIONS.items():
        for dst, route in row.items():
            if route.cloud:
                assert {formats.FORMATS[src].group, formats.FORMATS[dst].group} <= {
                    formats.DOCUMENT, formats.EBOOK, formats.SHEET}
                assert "pdf" in (src, dst) or formats.EBOOK in (formats.FORMATS[src].group,
                                                                 formats.FORMATS[dst].group)


def test_images_keep_local_engines_and_image_targets_exclude_pdf():
    assert formats.engine_for("png", "webp") == formats.E_IMAGE
    assert "pdf" not in formats.image_targets("png")
    assert "avif" in formats.targets("heic", formats.Capabilities())


def test_output_name_is_safe():
    assert formats.output_name("my report.final.docx", "pdf") == "my report.final.pdf"
    assert formats.output_name("../../etc/passwd", "txt").endswith(".txt")
    assert "/" not in formats.output_name("a/b/c.tar.gz", "zip")
    assert formats.output_name("a.tar.gz", "zip") == "a.zip"


def test_detect_prefers_bytes_over_name():
    assert detect.detect("pic.jpg", "image/jpeg", png()) == "png"
    assert detect.detect("data.json", None, b'{"a": 1}') == "json"
    assert detect.detect("blob", None, b"<?xml version='1.0'?><a/>") == "xml"
    assert detect.detect("notes.txt", None, b"hello") == "txt"
    assert detect.detect("x.unknown", None, b"\x00\x01\x02") is None


# --------------------------------------------------------------------------- #
# Image operations
# --------------------------------------------------------------------------- #


def test_resize_modes():
    plan = image_ops.plan_resize(400, 200, width=100, height=100, mode=image_ops.FIT)
    assert (plan.width, plan.height) == (100, 50)
    plan = image_ops.plan_resize(400, 200, width=100, height=100, mode=image_ops.FILL)
    assert (plan.width, plan.height) == (100, 100) and plan.crop is not None
    assert image_ops.plan_resize(400, 200, percent=50).width == 200
    with pytest.raises(MediaError) as err:
        image_ops.plan_resize(400, 200)
    assert err.value.code == "nothing_to_do"
    with pytest.raises(MediaError) as err:
        image_ops.plan_resize(400, 200, width=999999)
    assert err.value.code == "invalid_size"


def test_resize_compress_convert_roundtrip():
    raw = png((200, 100))
    done = image_ops.resize(raw, width=50)
    assert (done.width, done.height) == (50, 25) and done.format_key == "png"
    assert image_ops.convert(raw, "webp").format_key == "webp"
    flat = image_ops.convert(png((10, 10), (0, 0, 0, 0), "RGBA"), "jpg")
    assert "alpha_flattened" in flat.notes
    info = image_ops.read_info(raw)
    assert (info.format_key, info.width, info.height, info.has_alpha) == ("png", 200, 100, False)


def test_decompression_bomb_is_refused_from_the_header():
    big = io.BytesIO()
    Image.new("1", (9000, 9000)).save(big, "PNG")   # 81 MP, tiny file
    with pytest.raises(MediaError) as err:
        image_ops.read_info(big.getvalue())
    assert err.value.code == "too_many_pixels"
    with pytest.raises(MediaError) as err:
        image_ops.read_info(b"not an image")
    assert err.value.code == "unreadable"


# --------------------------------------------------------------------------- #
# Background removal (fake ONNX session)
# --------------------------------------------------------------------------- #


class FakeSession:
    def __init__(self, blob=True):
        self.blob = blob

    def get_inputs(self):
        return [types.SimpleNamespace(name="input.1")]

    def run(self, _outputs, feed):
        import numpy as np
        assert feed["input.1"].shape == (1, 3, 320, 320)
        out = np.zeros((1, 1, 320, 320), dtype="float32")
        if self.blob:
            out[:, :, 80:240, 80:240] = 1.0
        return [out]


def make_remover(tmp_path, **kw):
    remover = bg_removal.BackgroundRemover(tmp_path, **kw)
    remover.ensure_model = lambda fetch: asyncio.sleep(0)       # model file not needed here
    remover._load_session = lambda: setattr(remover, "_session", FakeSession()) or \
        setattr(remover, "_input_name", "input.1")
    return remover


async def test_background_removal_cuts_the_subject_out(tmp_path):
    remover = make_remover(tmp_path)
    raw = png((200, 120), (10, 200, 10))
    done = await remover.remove(raw, fetch=None)
    out = Image.open(io.BytesIO(done.data))
    assert done.format_key == "png" and out.mode == "RGBA" and out.size == (200, 120)
    assert out.getpixel((100, 60))[3] == 255          # centre kept
    assert out.getpixel((2, 2))[3] == 0               # corner transparent
    assert remover.loaded
    await remover.close()
    assert not remover.loaded


async def test_background_removal_without_subject_is_an_error(tmp_path):
    remover = make_remover(tmp_path)
    remover._load_session = lambda: setattr(remover, "_session", FakeSession(blob=False)) or \
        setattr(remover, "_input_name", "input.1")
    with pytest.raises(MediaError) as err:
        await remover.remove(png(), fetch=None)
    assert err.value.code == "no_subject"
    await remover.close()


async def test_model_unloads_after_idle(tmp_path):
    clock = {"now": 1000.0}
    remover = make_remover(tmp_path, idle_seconds=600, clock=lambda: clock["now"])
    await remover.remove(png(), fetch=None)
    assert remover.loaded
    clock["now"] += 599
    await asyncio.sleep(0.05)
    assert remover.loaded
    clock["now"] += 2
    for _ in range(100):
        await asyncio.sleep(0.05)
        if not remover.loaded:
            break
    assert not remover.loaded
    await remover.close()


async def test_inferences_never_overlap(tmp_path):
    remover = make_remover(tmp_path)
    running = {"now": 0, "max": 0}
    original = remover._infer_sync

    def spy(raw):
        running["now"] += 1
        running["max"] = max(running["max"], running["now"])
        try:
            return original(raw)
        finally:
            running["now"] -= 1

    remover._infer_sync = spy
    await asyncio.gather(*(remover.remove(png(), fetch=None) for _ in range(4)))
    assert running["max"] == 1
    await remover.close()


async def test_model_download_is_checksum_verified(tmp_path):
    payload = b"fake-onnx-bytes"
    good = bg_removal.sha256_of(payload)
    calls = []

    async def fetch(url):
        calls.append(url)
        return payload

    remover = bg_removal.BackgroundRemover(tmp_path, sha256=good)
    path = await remover.ensure_model(fetch)
    assert path.read_bytes() == payload and len(calls) == 1
    await remover.ensure_model(fetch)               # already on disk and valid
    assert len(calls) == 1

    bad = bg_removal.BackgroundRemover(tmp_path / "other", sha256="0" * 64)
    with pytest.raises(MediaError) as err:
        await bad.ensure_model(fetch)
    assert err.value.code == "model_unavailable"
    assert not bad.model_on_disk()


def test_official_model_constants_look_right():
    assert re.fullmatch(r"[0-9a-f]{64}", bg_removal.MODEL_SHA256)
    assert bg_removal.MODEL_URL.startswith("https://github.com/danielgatis/rembg/")
    assert bg_removal.IDLE_SECONDS == 600


# --------------------------------------------------------------------------- #
# Local converters
# --------------------------------------------------------------------------- #


def conv(src, dst, raw, name="file"):
    return local.convert_bytes(src, dst, raw, name)


def test_data_conversions():
    data = b'{"a": [{"x": 1, "y": "h\\u00e9"}, {"x": 2, "y": null}], "n": null}'
    assert b"x: 1" in conv("json", "yaml", data).data
    toml = conv("json", "toml", data)
    assert "null_dropped" in toml.notes and b"null" not in toml.data
    assert b"<root>" in conv("json", "xml", data).data
    assert json.loads(conv("yaml", "json", conv("json", "yaml", data).data).data)["a"][0]["x"] == 1
    assert conv("json", "csv", b'[{"a":1,"b":2}]').data == b"a,b\n1,2\n"
    with pytest.raises(MediaError) as err:
        conv("json", "json", b"{broken")
    assert err.value.code in ("unsupported", "invalid_data")
    with pytest.raises(MediaError) as err:
        conv("json", "yaml", b"{broken")
    assert err.value.code == "invalid_data"
    with pytest.raises(MediaError) as err:
        conv("json", "csv", b'"just a string"')
    assert err.value.code == "not_tabular"


def test_yaml_alias_bomb_is_refused():
    bomb = ("a: &a [1,2,3]\n" + "".join(f"k{i}: *a\n" for i in range(150))).encode()
    with pytest.raises(MediaError) as err:
        conv("yaml", "json", bomb)
    assert err.value.code == "unsafe_data"


def test_spreadsheet_roundtrips_and_text_stays_text():
    rows = b"name,note\nAda,=1+1\nBob,hello\n"
    xlsx = conv("csv", "xlsx", rows)
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(xlsx.data)).active
    assert ws["B2"].value == "=1+1" and ws["B2"].data_type == "s"       # never a formula
    assert conv("xlsx", "csv", xlsx.data).data.startswith(b"\xef\xbb\xbfname,note")
    ods = conv("csv", "ods", rows)
    assert detect.detect("x.ods", None, ods.data) == "ods"
    assert b"Ada" in conv("ods", "csv", ods.data).data
    recs = json.loads(conv("xlsx", "json", xlsx.data).data)
    assert recs[1] == {"name": "Bob", "note": "hello"}


def test_too_many_cells_are_refused(monkeypatch):
    monkeypatch.setattr(formats, "MAX_SHEET_CELLS", 10)
    with pytest.raises(MediaError) as err:
        conv("csv", "xlsx", ("a,b,c\n" * 10).encode())
    assert err.value.code == "too_many_cells"


def test_document_conversions():
    html = conv("md", "html", b"# Title\n\n* one\n* two", "note.md").data.decode()
    assert "<h1>Title</h1>" in html and "<li>one</li>" in html and "charset" in html
    md = conv("html", "md", b"<h1>T</h1><p>a <b>b</b></p><script>alert(1)</script>").data.decode()
    assert md.startswith("# T") and "alert" not in md
    assert conv("html", "txt", b"<p>one</p><p>two</p>").data.decode().split() == ["one", "two"]
    assert conv("txt", "txt", "caf\xe9".encode("cp1252")).data.decode("utf-8") == "café"


def make_docx() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml",
                   '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr("_rels/.rels",
                   '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr("word/document.xml",
                   '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                   "<w:body><w:p><w:r><w:t>Hello docx</w:t></w:r></w:p></w:body></w:document>")
    return buf.getvalue()


def test_docx_to_text_markdown_html():
    raw = make_docx()
    assert detect.detect("a.docx", None, raw) == "docx"
    assert conv("docx", "txt", raw).data.decode().strip() == "Hello docx"
    assert "Hello docx" in conv("docx", "md", raw).data.decode()
    assert "<p>Hello docx</p>" in conv("docx", "html", raw).data.decode()


def test_pdf_roundtrip_and_text():
    pdf = conv("png", "pdf", png((300, 200)), "scan.png")
    assert pdf.data.startswith(b"%PDF")
    page = conv("pdf", "png", pdf.data, "scan.pdf")
    assert page.name == "scan.png" and page.pages == 1 and page.format_key == "png"
    with pytest.raises(MediaError) as err:
        conv("pdf", "txt", pdf.data)
    assert err.value.code == "no_text"


def test_archives_convert_and_unsafe_members_are_skipped():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("dir/ok.txt", "hello")
        z.writestr("../evil.txt", "x")
        z.writestr("/abs.txt", "x")
    for dst in ("tgz", "7z"):
        out = conv("zip", dst, buf.getvalue(), "pack.zip")
        assert "entries_skipped" in out.notes
        assert detect.detect(out.name, None, out.data) == dst
    back = conv("7z", "zip", conv("zip", "7z", buf.getvalue()).data)
    with zipfile.ZipFile(io.BytesIO(back.data)) as z:
        assert z.namelist() == ["dir/ok.txt"]


def test_zip_bomb_ratio_and_file_count(monkeypatch):
    monkeypatch.setattr(formats, "MAX_ARCHIVE_FILES", 3)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for i in range(5):
            z.writestr(f"{i}.txt", "x")
    with pytest.raises(MediaError) as err:
        conv("zip", "tgz", buf.getvalue())
    assert err.value.code == "too_many_files"
    monkeypatch.setattr(formats, "MAX_ARCHIVE_FILES", 1000)
    big = io.BytesIO()
    with zipfile.ZipFile(big, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("zeros.bin", b"\0" * (30 * 1024 * 1024))
    with pytest.raises(MediaError) as err:
        conv("zip", "tgz", big.getvalue())
    assert err.value.code == "unsafe_archive"


def test_safe_member_names():
    assert local.safe_member_name("a/b.txt") == "a/b.txt"
    assert local.safe_member_name("./a//b") == "a/b"
    for bad in ("../x", "a/../../x", "/etc/passwd", "C:\\x", ""):
        assert local.safe_member_name(bad) is None


def test_unsupported_and_oversized_inputs():
    with pytest.raises(MediaError) as err:
        conv("png", "mp3", png())
    assert err.value.code == "unsupported"
    with pytest.raises(MediaError) as err:
        conv("docx", "pdf", make_docx())          # cloud only: never run locally
    assert err.value.code == "unsupported"
    with pytest.raises(MediaError) as err:
        conv("json", "yaml", b" " * (formats.MAX_INPUT_BYTES + 1))
    assert err.value.code == "too_large"


# --------------------------------------------------------------------------- #
# ffmpeg
# --------------------------------------------------------------------------- #


def test_ffmpeg_command_is_locked_down():
    cmd = media.build_command("ffmpeg", "mp4", "webm", Path("/t/in.mp4"), Path("/t/out.webm"), 60)
    assert cmd[cmd.index("-protocol_whitelist") + 1] == "file"
    assert cmd[cmd.index("-f") + 1] == "mov"           # demuxer forced from the sniffed format
    assert "-nostdin" in cmd and "-t" in cmd and "-fs" in cmd
    assert cmd[-1] == "/t/out.webm"
    gif = media.build_command("ffmpeg", "mp4", "gif", Path("in.mp4"), Path("out.gif"), 20)
    assert "-filter_complex" in gif and "[out]" in gif
    with pytest.raises(MediaError):
        media.build_command("ffmpeg", "mp3", "png", Path("a"), Path("b"), 10)


def test_ffmpeg_duration_caps_and_parsing():
    assert media.parse_duration("  Duration: 00:01:02.50, start: 0") == 62.5
    assert media.parse_duration("nothing") is None
    assert media.max_seconds("mp3", "wav") == formats.MAX_AUDIO_SECONDS
    assert media.max_seconds("mp4", "mp3") == formats.MAX_AUDIO_SECONDS
    assert media.max_seconds("mp4", "webm") == formats.MAX_VIDEO_SECONDS
    assert media.max_seconds("mp4", "gif") == formats.MAX_GIF_SECONDS


@pytest.mark.skipif(not media.available(), reason="imageio-ffmpeg binary not installed")
async def test_ffmpeg_really_converts_a_tiny_wav():
    import struct
    import wave
    buf = io.BytesIO()
    with wave.open(buf, "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(8000)
        w.writeframes(b"".join(struct.pack("<h", (i % 100) * 100) for i in range(8000)))
    out = await media.convert_media("wav", "mp3", buf.getvalue(), "beep.wav")
    assert out.name == "beep.mp3" and detect.detect(out.name, None, out.data) == "mp3"
    with pytest.raises(MediaError) as err:
        await media.convert_media("mp3", "wav", b"ID3" + b"\0" * 64, "x.mp3")
    assert err.value.code in ("unreadable", "convert_failed", "no_stream")


# --------------------------------------------------------------------------- #
# CloudConvert adapter + service mapping
# --------------------------------------------------------------------------- #


def test_cloudconvert_job_and_result_parsing():
    job = cc.build_job("https://cdn.discordapp.com/x", "docx", "pdf", "input.docx", "r.pdf")
    assert job["tasks"]["import"]["url"].startswith("https://cdn.discordapp.com")
    assert job["tasks"]["convert"]["filename"] == "r.pdf"
    finished = {"data": {"status": "finished", "tasks": [
        {"operation": "export/url", "status": "finished",
         "result": {"files": [{"url": "https://storage.cloudconvert.com/a.pdf",
                               "filename": "a.pdf", "size": 1234}]}}]}}
    assert cc.parse_finished(finished) == {"url": "https://storage.cloudconvert.com/a.pdf",
                                           "filename": "a.pdf", "size": 1234}
    failed = {"data": {"status": "error", "tasks": [{"status": "error", "code": "CONVERSION_FAILED"}]}}
    with pytest.raises(ProviderError) as err:
        cc.parse_finished(failed)
    assert err.value.status == 422 and "CONVERSION_FAILED" in err.value.body


def test_cloudconvert_adapter_needs_a_key_and_hides_the_url_from_payload():
    with pytest.raises(ConfigurationError):
        cc.CloudConvertAdapter(None)
    from gateway.clients.conversion import CloudConvertClient
    seen = {}

    class Executor:
        async def execute(self, spec):
            seen["spec"] = spec
            return {"url": "https://storage.cloudconvert.com/a.pdf", "size": 1}

    asyncio.run(CloudConvertClient(Executor()).convert(
        import_url="https://cdn.discordapp.com/secret", input_format="docx", output_format="pdf",
        filename="input.docx", quota=[QuotaTarget.user(1, "convert_cloud")]))
    spec = seen["spec"]
    assert spec.provider == "cloudconvert" and spec.operation == "convert"
    assert "secret" not in json.dumps(spec.payload)             # payload is logged to staff
    assert spec.metadata["import_url"].endswith("secret")


async def test_cloudconvert_adapter_polls_until_finished(monkeypatch):
    posts, gets = [], []

    class Resp:
        def __init__(self, data):
            self._data, self.status, self.headers = data, 200, {}

        async def __aenter__(self):
            return self

        async def __aexit__(self, *a):
            return False

        async def json(self, content_type=None):
            return self._data

        async def text(self):
            return ""

    class Session:
        def request(self, method, url, **kw):
            if method == "POST":
                posts.append(kw["json"])
                return Resp({"data": {"id": "job1", "status": "waiting", "tasks": []}})
            gets.append(url)
            status = "finished" if len(gets) > 1 else "processing"
            tasks = [{"operation": "export/url", "status": "finished", "result": {"files": [
                {"url": "https://storage.cloudconvert.com/x.pdf", "filename": "x.pdf", "size": 9}]}}]
            return Resp({"data": {"id": "job1", "status": status, "tasks": tasks}})

    monkeypatch.setattr(cc, "_POLL_SECONDS", 0)
    adapter = cc.CloudConvertAdapter("key")
    adapter._session = Session()
    from gateway.spec import CallSpec
    spec = CallSpec(provider="cloudconvert", operation="convert", model="convert",
                    payload={"input_format": "docx", "output_format": "pdf", "filename": "input.docx"},
                    quota=[], call_type="convert_cloud", correlation_id="c",
                    metadata={"import_url": "https://cdn.discordapp.com/a"})
    result = await adapter.execute(spec)
    assert result.data["url"].endswith("x.pdf") and len(gets) == 2 and len(posts) == 1


def test_gateway_errors_map_to_stable_codes():
    assert svc.gateway_error(QuotaExceededError(QuotaTarget.user(1, "convert_cloud"))).code == "quota"
    assert svc.gateway_error(ModelRateLimitError("cloudconvert", "convert", "rpd", 25)).code == "cloud_busy"
    assert svc.gateway_error(RateLimitError("cloudconvert")).code == "cloud_busy"
    assert svc.gateway_error(ProviderError("cloudconvert", 422, "CONVERSION_FAILED")).code == "convert_failed"
    assert svc.gateway_error(ProviderError("cloudconvert", 402, "credits")).code == "cloud_unavailable"
    assert svc.gateway_error(APIUnavailableError("cloudconvert")).code == "cloud_unavailable"


def fake_bot(cloud=False, job=None, body=b"%PDF-1.4 converted", error=None):
    calls = []

    class Cloud:
        async def convert(self, **kw):
            calls.append(kw)
            if error:
                raise error
            return job or {"url": "https://storage.cloudconvert.com/o.pdf", "size": len(body)}

    class Api:
        async def request(self, method, url, **kw):
            return types.SimpleNamespace(body=body)

    gw = types.SimpleNamespace(cloudconvert=Cloud() if cloud else None,
                               cloudconvert_available=lambda: cloud)
    return types.SimpleNamespace(gateway=gw, public_api=Api(), redis=None), calls


async def test_cloud_conversion_quota_target_and_reupload():
    bot, calls = fake_bot(cloud=True)
    service = svc.MediaToolsService(bot)
    result = await service.convert("docx", "pdf", b"x" * 10, "report.docx", user_id=7, guild_id=9,
                                   source_url="https://cdn.discordapp.com/a", upload_limit=10**7)
    assert result.cloud and result.output.data.startswith(b"%PDF") and result.output.name == "report.pdf"
    quota = calls[0]["quota"]
    assert [(q.scope.value, q.type) for q in quota] == [
        ("user", "convert_cloud"), ("guild", "convert_cloud"), ("global", "convert_cloud")]


async def test_cloud_conversion_over_upload_limit_returns_a_link():
    bot, _ = fake_bot(cloud=True, job={"url": "https://storage.cloudconvert.com/o.pdf", "size": 50_000_000})
    result = await svc.MediaToolsService(bot).convert(
        "docx", "pdf", b"x", "r.docx", user_id=1, guild_id=None,
        source_url="https://cdn.discordapp.com/a", upload_limit=10**7)
    assert result.output is None and result.link.endswith("o.pdf") and result.link_size == 50_000_000


async def test_cloud_link_from_a_foreign_host_is_refused_and_missing_key_hides_routes():
    bot, _ = fake_bot(cloud=True, job={"url": "https://evil.example/o.pdf", "size": 1})
    with pytest.raises(MediaError) as err:
        await svc.MediaToolsService(bot).convert("docx", "pdf", b"x", "r.docx", user_id=1,
                                                 guild_id=None, source_url="https://cdn.discordapp.com/a",
                                                 upload_limit=10**7)
    assert err.value.code == "convert_failed"
    nokey, _ = fake_bot(cloud=False)
    service = svc.MediaToolsService(nokey)
    assert "pdf" not in service.targets("docx")
    with pytest.raises(MediaError) as err:
        await service.convert("docx", "pdf", b"x", "r.docx", user_id=1, guild_id=None,
                              source_url="https://cdn.discordapp.com/a", upload_limit=10**7)
    assert err.value.code == "unsupported"


async def test_local_conversion_through_the_service():
    bot, _ = fake_bot()
    service = svc.MediaToolsService(bot)
    result = await service.convert("json", "yaml", b'{"a": 1}', "d.json", user_id=1, guild_id=None,
                                   source_url=None, upload_limit=10**7)
    assert result.output.data == b"a: 1\n" and not result.cloud
    assert service.identify("d.json", None, b'{"a": 1}') == "json"
    with pytest.raises(MediaError) as err:
        service.identify("thing.xyz", None, b"\x00\x01\x02")
    assert err.value.code == "unsupported_type"


# --------------------------------------------------------------------------- #
# Cards and persistent components
# --------------------------------------------------------------------------- #


def test_resolve_format_accepts_labels_extensions_and_dots():
    assert views.resolve_format("PDF") == "pdf"
    assert views.resolve_format(".jpeg") == "jpg"
    assert views.resolve_format("Markdown") == "md"
    assert views.resolve_format("nope") is None


def test_image_card_shows_before_after_and_the_four_buttons():
    before = image_ops.read_info(png((400, 300)))
    after = image_ops.read_info(image_ops.convert(png((400, 300)), "webp").data)
    files = views.CardFiles([], "a.webp", None)
    card = views.build_image_card(after, "en-US", files, before=before, notes=["animation_lost"])
    text = walk_text(card)
    assert "PNG → WebP" in text and "400 × 300 px" in text and "→" in text
    ids = custom_ids(card)
    assert ids == ["moddy:tools:image:resize", "moddy:tools:image:compress",
                   "moddy:tools:image:convert", "moddy:tools:image:bg"]
    assert "only the first frame was kept" in text
    # a card without its source (above the upload limit) carries no buttons
    assert custom_ids(views.build_image_card(after, "fr", files, buttons=False)) == []


def test_non_previewable_formats_show_a_png_preview_and_the_original_as_a_file():
    class Service:
        async def preview_png(self, data):
            return png((8, 8))

    heic_like = image_ops.ImageInfo("tiff", 10, 10, 100, "RGB", False, False, 1)
    files = asyncio.run(views.make_card_files(Service(), png(), "pic.tiff", heic_like, 10**7))
    assert [f.filename for f in files.files] == ["pic.tiff", views.PREVIEW_NAME]
    assert files.gallery == views.PREVIEW_NAME and files.file_ref == "pic.tiff"
    over = asyncio.run(views.make_card_files(Service(), png(), "pic.tiff", heic_like, 5))
    assert [f.filename for f in over.files] == [views.PREVIEW_NAME] and over.file_ref is None


def test_find_source_attachment_skips_the_preview():
    att = lambda name, ctype="image/png": types.SimpleNamespace(filename=name, content_type=ctype)
    msg = types.SimpleNamespace(attachments=[att(views.PREVIEW_NAME), att("photo.tiff", "image/tiff")])
    assert views.find_source_attachment(msg).filename == "photo.tiff"
    assert views.find_source_attachment(types.SimpleNamespace(attachments=[att(views.PREVIEW_NAME)])) is None


def _v2_message(*components):
    """A Components V2 card as Discord returns it: no attachments, only components."""
    from discord.components import _component_factory
    container = _component_factory({"type": 17, "id": 1, "components": list(components)}, None)
    return types.SimpleNamespace(attachments=[], components=[container])


def _gallery(name, ctype="image/png"):
    return {"type": 12, "id": 2, "items": [{"media": {
        "url": f"https://cdn.discordapp.com/attachments/1/2/{name}?ex=1", "content_type": ctype}}]}


def test_image_is_found_in_the_components_of_a_v2_card():
    message = _v2_message(_gallery("photo.png"))
    found = views.find_source_attachment(message)
    assert found.filename == "photo.png" and found.url.startswith("https://cdn.discordapp.com/")
    assert views.is_image_attachment(found)


def test_v2_card_with_a_preview_returns_the_original_file_not_the_preview():
    message = _v2_message(
        _gallery(views.PREVIEW_NAME),
        {"type": 13, "id": 3, "name": "scan.tiff", "size": 99,
         "file": {"url": "https://cdn.discordapp.com/attachments/1/3/scan.tiff?ex=1"}})
    assert views.find_source_attachment(message).filename == "scan.tiff"
    only_preview = _v2_message(_gallery(views.PREVIEW_NAME))
    assert views.find_source_attachment(only_preview) is None
    info = image_ops.ImageInfo("tiff", 10, 10, 100, "RGB", False, False, 1)
    files = views.card_files_from_message(message, views.find_source_attachment(message), info)
    assert files.gallery == views.PREVIEW_NAME and files.file_ref == "scan.tiff"


async def test_card_image_reads_through_its_media_item():
    class Media:
        url = "https://cdn.discordapp.com/attachments/1/2/a%20b.png?ex=1"
        content_type = "image/png"

        async def read(self):
            return b"bytes"

    card = views.CardImage(Media())
    assert card.filename == "a b.png" and await card.read() == b"bytes"


async def test_card_image_is_fetched_back_when_the_interaction_message_has_none():
    att = lambda name: types.SimpleNamespace(filename=name, content_type="image/png")
    bare = types.SimpleNamespace(id=1, attachments=[], flags=discord.MessageFlags._from_value(0))
    full = types.SimpleNamespace(attachments=[att("photo.png")])

    class Interaction:
        message = bare

        async def original_response(self):
            return full

    assert (await views.find_card_attachment(Interaction())).filename == "photo.png"

    class Gone(Interaction):
        async def original_response(self):
            return bare

    assert await views.find_card_attachment(Gone()) is None


async def test_modals_keep_the_message_the_button_was_clicked_on():
    card = types.SimpleNamespace(attachments=[types.SimpleNamespace(filename="a.png", content_type="image/png")])
    stripped = types.SimpleNamespace(attachments=[])
    interaction = types.SimpleNamespace(message=stripped)
    modal = views.CompressModal("en-US", card)
    seen = {}

    async def fake_apply(inter, op, params):
        seen["message"] = inter.message

    original = views.apply_image_action
    views.apply_image_action = fake_apply
    try:
        await modal.on_submit(interaction)
    finally:
        views.apply_image_action = original
    assert seen["message"] is card


def test_convert_picker_groups_targets_and_flags_cloud_formats():
    targets = ["png", "jpg", "pdf", "docx"]
    view = views.build_convert_picker(filename="x.docx", src="docx", size=2048, token="abcdefghij",
                                      targets=["pdf", "md", "txt", "html"], locale="en-US")
    ids = custom_ids(view)
    assert ids == ["moddy:tools:convert:to:abcdefghij:document"]
    select = next(i.item for i in view.walk_children() if isinstance(getattr(i, 'item', None), ui.Select))
    cloud_option = next(o for o in select.options if o.value == "pdf")
    assert cloud_option.description and "5" in cloud_option.description
    assert views.grouped_targets(targets) == {"image": ["png", "jpg"], "document": ["pdf", "docx"]}
    assert "CloudConvert" in walk_text(view)


def test_dynamic_items_roundtrip_and_do_not_collide():
    seen = set()
    for cls, args in views.DYNAMIC_ITEM_CASES:
        cid = cls(*args).item.custom_id
        assert cls.__discord_ui_compiled_template__.fullmatch(cid) and len(cid) <= 100
        assert cid not in seen
        seen.add(cid)
    token = asyncio.run(tc.save_state(types.SimpleNamespace(redis=None), "convert", {"k": "png"}))
    assert re.fullmatch(r"[A-Za-z0-9]{10}", token)
    select = views.ConvertSelect(token, "image", "en-US", src="png", keys=["jpg", "webp"])
    match = views.ConvertSelect.__discord_ui_compiled_template__.fullmatch(select.item.custom_id)
    assert match["token"] == token and match["group"] == "image"


# --------------------------------------------------------------------------- #
# Flows with a fake interaction (no Discord)
# --------------------------------------------------------------------------- #


class FakeResponse:
    def __init__(self):
        self.done, self.calls = False, []

    def is_done(self):
        return self.done

    async def defer(self, **kw):
        self.done = True
        self.calls.append(("defer", kw))

    async def edit_message(self, **kw):
        self.done = True
        self.calls.append(("edit_message", kw))

    async def send_message(self, **kw):
        self.done = True
        self.calls.append(("send_message", kw))

    async def send_modal(self, modal):
        self.done = True
        self.calls.append(("send_modal", modal))


class FakeAttachment:
    def __init__(self, data, filename="a.png", content_type="image/png"):
        self._data, self.filename, self.content_type = data, filename, content_type
        self.size, self.url = len(data), f"https://cdn.discordapp.com/attachments/1/2/{filename}"

    async def read(self):
        return self._data


class FakeInteraction:
    def __init__(self, bot, *, message=None, user_id=5, limit=10**7):
        self.client, self.message, self.guild_id = bot, message, None
        self.user = types.SimpleNamespace(id=user_id)
        self.locale, self.filesize_limit = "en-US", limit
        self.response = FakeResponse()
        self.edits, self.followups = [], []
        self.followup = types.SimpleNamespace(send=self._followup)

    async def edit_original_response(self, **kw):
        self.edits.append(kw)

    async def _followup(self, **kw):
        self.followups.append(kw)


def flow_bot():
    bot, _ = fake_bot()
    bot.media_tools = svc.MediaToolsService(bot)
    return bot


async def test_image_command_builds_the_card_and_attaches_the_image():
    bot = flow_bot()
    interaction = FakeInteraction(bot)
    await views.start_image(interaction, FakeAttachment(png((300, 200))), ephemeral=True)
    assert interaction.response.calls[0][0] == "defer"
    edit = interaction.edits[-1]
    assert [f.filename for f in edit["attachments"]] == ["a.png"]
    text = walk_text(edit["view"])
    assert "300 × 200 px" in text and "PNG" in text
    assert len(custom_ids(edit["view"])) == 4


async def test_image_command_refuses_a_non_image():
    bot = flow_bot()
    interaction = FakeInteraction(bot)
    await views.start_image(interaction, FakeAttachment(b"x", "doc.pdf", "application/pdf"),
                            ephemeral=True)
    assert "not an image" in walk_text(interaction.edits[-1]["view"])


async def test_resize_button_flow_edits_the_card_in_place_with_before_after():
    bot = flow_bot()
    source = FakeAttachment(png((400, 200)), "a.png")
    message = types.SimpleNamespace(attachments=[source],
                                    flags=types.SimpleNamespace(ephemeral=True),
                                    interaction_metadata=None)
    interaction = FakeInteraction(bot, message=message)
    await views.apply_image_action(interaction, "resize", {"width": 100, "height": None,
                                                           "percent": None, "mode": "fit"})
    assert interaction.response.calls[0][0] == "edit_message"       # loading card, in place
    edit = interaction.edits[-1]
    text = walk_text(edit["view"])
    assert "400 × 200 px → 100 × 50 px" in text
    assert len(custom_ids(edit["view"])) == 4                       # same buttons again


async def test_public_card_clicked_by_someone_else_answers_privately():
    bot = flow_bot()
    message = types.SimpleNamespace(
        attachments=[FakeAttachment(png())], flags=types.SimpleNamespace(ephemeral=False),
        interaction_metadata=types.SimpleNamespace(user=types.SimpleNamespace(id=1)))
    interaction = FakeInteraction(bot, message=message, user_id=2)
    assert not views.owns_card(interaction)
    await views.apply_image_action(interaction, "compress", {"level": "light"})
    assert interaction.response.calls[0][0] == "send_message"
    assert interaction.response.calls[0][1]["ephemeral"] is True


async def test_failed_action_keeps_the_image_card_and_reports_the_error():
    bot = flow_bot()
    message = types.SimpleNamespace(attachments=[FakeAttachment(png((50, 50)), "a.png")],
                                    flags=types.SimpleNamespace(ephemeral=True),
                                    interaction_metadata=None)
    interaction = FakeInteraction(bot, message=message)
    await views.apply_image_action(interaction, "resize", {"width": 99999, "height": None,
                                                           "percent": None, "mode": "fit"})
    assert len(custom_ids(interaction.edits[-1]["view"])) == 4       # card restored
    assert interaction.followups and "between 1 and" in walk_text(interaction.followups[-1]["view"])


async def test_result_above_the_upload_limit_is_an_error_not_a_crash():
    bot = flow_bot()
    message = types.SimpleNamespace(attachments=[FakeAttachment(png((200, 200)), "a.png")],
                                    flags=types.SimpleNamespace(ephemeral=True),
                                    interaction_metadata=None)
    interaction = FakeInteraction(bot, message=message, limit=500)
    await views.apply_image_action(interaction, "resize", {"width": 2000, "height": None,
                                                           "percent": None, "mode": "fit"})
    assert "upload limit" in walk_text((interaction.followups or interaction.edits)[-1]["view"])


async def test_convert_command_lists_only_compatible_formats_then_converts():
    bot = flow_bot()
    state_store = {}

    async def save(bot_, kind, payload, **kw):
        state_store["x"] = payload
        return "abcdefghij"

    original = tc.save_state
    tc.save_state = save
    try:
        interaction = FakeInteraction(bot)
        await views.start_convert(interaction, FakeAttachment(b'{"a": 1}', "d.json", "application/json"),
                                  None, ephemeral=True)
    finally:
        tc.save_state = original
    view = interaction.edits[-1]["view"]
    ids = custom_ids(view)
    assert any(i.endswith(":data") for i in ids) and not any(i.endswith(":audio") for i in ids)
    assert state_store["x"]["k"] == "json"

    # direct conversion with the format option
    async def read_url(url, limit=0):
        return b'{"a": 1}'

    bot.media_tools.read_url = read_url
    direct = FakeInteraction(bot)
    await views.start_convert(direct, FakeAttachment(b'{"a": 1}', "d.json", "application/json"),
                              "yaml", ephemeral=True)
    result = direct.edits[-1]
    assert [f.filename for f in result["attachments"]] == ["d.yaml"]
    assert "JSON → YAML" in walk_text(result["view"])

    bad = FakeInteraction(bot)
    await views.start_convert(bad, FakeAttachment(b'{"a": 1}', "d.json", "application/json"),
                              "mp3", ephemeral=True)
    assert "cannot become" in walk_text(bad.edits[-1]["view"])


async def test_unknown_file_type_is_explained():
    bot = flow_bot()
    interaction = FakeInteraction(bot)
    await views.start_convert(interaction, FakeAttachment(b"\x00\x01\x02", "mystery.xyz", None),
                              None, ephemeral=True)
    assert "not supported" in walk_text(interaction.edits[-1]["view"])


# --------------------------------------------------------------------------- #
# i18n
# --------------------------------------------------------------------------- #


def flatten(node, prefix=""):
    out = {}
    for key, value in node.items():
        path = f"{prefix}.{key}" if prefix else key
        if isinstance(value, dict):
            out.update(flatten(value, path))
        else:
            out[path] = value
    return out


def locale_keys(locale, root):
    with open(ROOT / "locales" / f"{locale}.json", encoding="utf-8") as f:
        return flatten(json.load(f)["tools"][root], f"tools.{root}")


@pytest.mark.parametrize("root", ["image", "convert"])
def test_keys_and_placeholders_match_in_the_five_locales(root):
    reference = locale_keys("en-US", root)
    for locale in LOCALES:
        keys = locale_keys(locale, root)
        assert set(keys) == set(reference), (locale, set(keys) ^ set(reference))
        for key, text in keys.items():
            assert text.strip(), (locale, key)
            assert sorted(re.findall(r"\{(\w+)\}", text)) == \
                sorted(re.findall(r"\{(\w+)\}", reference[key])), (locale, key)


def test_every_error_code_raised_in_code_has_a_message():
    sources = [ROOT / "tools" / "image_ops.py", ROOT / "tools" / "bg_removal.py",
               ROOT / "tools" / "convert" / "local.py", ROOT / "tools" / "convert" / "media.py",
               ROOT / "services" / "media_tools_service.py", ROOT / "utils" / "tools_media_views.py"]
    codes = set()
    for path in sources:
        codes |= set(re.findall(r'MediaError\(\s*"(\w+)"', path.read_text(encoding="utf-8")))
    convert_errors = {k.rsplit(".", 1)[1] for k in locale_keys("en-US", "convert")
                      if ".errors." in k}
    image_errors = {k.rsplit(".", 1)[1] for k in locale_keys("en-US", "image") if ".errors." in k}
    shared = {"unsupported_type", "bad_target"}
    for code in codes:
        assert code in convert_errors or code in image_errors, code
    # the image flow only ever raises the image codes + the ones of the shared converters
    assert {"not_image", "invalid_size", "no_subject", "model_unavailable"} <= image_errors
    for code in codes - shared:
        if code not in image_errors:
            assert code in convert_errors


def test_card_texts_resolve_in_every_locale():
    info = image_ops.read_info(png())
    for locale in LOCALES:
        card = views.build_image_card(info, locale, views.CardFiles([], "a.png", None),
                                      before=info, notes=["animation_lost"])
        text = walk_text(card)
        assert not re.search(r"\btools\.(image|convert)\.", text), (locale, text)
        picker = views.build_convert_picker(filename="a.png", src="png", size=10, token="abcdefghij",
                                            targets=["jpg", "pdf"], locale=locale)
        assert not re.search(r"\btools\.(image|convert)\.", walk_text(picker)), locale


@pytest.mark.parametrize("locale", LOCALES)
async def test_modals_build_within_discord_limits(locale):
    for modal in (views.ResizeModal(locale), views.CompressModal(locale)):
        assert len(modal.title) <= 45
        top = [c for c in modal.children]
        assert 1 <= len(top) <= 5
        for item in top:
            if isinstance(item, ui.Label):
                assert len(item.text) <= 45 and len(item.description or "") <= 100
    resize = views.ResizeModal(locale)
    assert [o.value for o in resize.mode.component.options] == ["fit", "exact", "fill"]
    assert sum(o.default for o in resize.mode.component.options) == 1
    compress = views.CompressModal(locale)
    assert [o.value for o in compress.level.component.options] == ["light", "balanced", "strong"]
